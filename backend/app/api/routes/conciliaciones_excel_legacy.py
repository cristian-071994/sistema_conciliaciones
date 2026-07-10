from io import BytesIO
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from app.models.conciliacion import Conciliacion
from app.models.conciliacion_item import ConciliacionItem
from app.models.enums import ItemTipo, UserRole
from app.models.viaje import Viaje
from .conciliaciones_helpers import (
    _extract_liquidacion_metadata,
    _normalize_manifiesto_for_lookup,
    _as_float,
    _split_total_evenly,
    _display_estado,
    _item_servicio_codigo,
    _fetch_avansat_with_fallback,
)


def _prepare_facturacion_rows(
    db: Session,
    items: list[ConciliacionItem],
    avansat_prefetched: dict[str, dict] | None = None,
) -> tuple[list[dict], list[str]]:
    rows: list[dict] = []
    missing_manifiestos: list[str] = []

    for item in items:
        viaje_ref = f"viaje #{item.viaje_id}" if item.viaje_id else f"item #{item.id}"
        manifiesto = _normalize_manifiesto_for_lookup(item.manifiesto_numero)
        if not manifiesto:
            missing_manifiestos.append(f"{viaje_ref} (sin manifiesto)")
            continue

        avansat = _fetch_avansat_with_fallback(db, manifiesto, avansat_prefetched)
        if not avansat:
            missing_manifiestos.append(f"{viaje_ref} (manifiesto {manifiesto} sin datos en Avansat)")
            continue

        precio_cliente = _as_float(item.tarifa_cliente)
        precio_tercero = _as_float(item.tarifa_tercero)
        rentabilidad = _as_float(item.rentabilidad)
        ganancia = precio_cliente - precio_tercero

        rows.append(
            {
                "manifiesto": manifiesto,
                "fecha_emision": avansat.get("fecha_emision") or "",
                "producto": avansat.get("producto") or "",
                "placa_vehiculo": avansat.get("placa_vehiculo") or (item.placa or ""),
                "trayler": avansat.get("trayler") or "",
                "remesa": avansat.get("remesa") or (item.remesa or ""),
                "ciudad_origen": avansat.get("ciudad_origen") or (item.origen or ""),
                "ciudad_destino": avansat.get("ciudad_destino") or (item.destino or ""),
                "precio_cliente": precio_cliente,
                "precio_tercero": precio_tercero,
                "rentabilidad": rentabilidad,
                "ganancia": ganancia,
            }
        )

    return rows, sorted(set(missing_manifiestos))


def _build_facturacion_excel(conc: Conciliacion, rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Facturacion"

    header_fill = PatternFill(fill_type="solid", fgColor="E5E7EB")
    total_fill = PatternFill(fill_type="solid", fgColor="FFF200")
    header_font = Font(bold=True, color="1F2937")
    total_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    cop_format = '"$" #,##0'

    ws.append(
        [
            "Manifiesto",
            "Fecha Manifiesto",
            "Producto",
            "Placa",
            "Remolque",
            "Remesa",
            "Origen",
            "Destino",
            "Precio Cliente",
            "Precio Tercero",
            "Rentabilidad",
            "Ganancia Cointra",
        ]
    )

    for idx, cell in enumerate(ws[1], start=1):
        cell.fill = header_fill
        if idx in (9, 10, 11, 12):
            cell.fill = total_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    total_precio_cliente = 0.0
    total_precio_tercero = 0.0
    total_ganancia = 0.0

    for row in rows:
        manifiesto = str(row.get("manifiesto") or "").strip()
        precio_cliente = _as_float(row.get("precio_cliente"))
        precio_tercero = _as_float(row.get("precio_tercero"))
        rentabilidad = _as_float(row.get("rentabilidad"))
        ganancia = _as_float(row.get("ganancia"))

        total_precio_cliente += precio_cliente
        total_precio_tercero += precio_tercero
        total_ganancia += ganancia

        ws.append(
            [
                manifiesto,
                row.get("fecha_emision") or "",
                row.get("producto") or "",
                row.get("placa_vehiculo") or "",
                row.get("trayler") or "",
                row.get("remesa") or "",
                row.get("ciudad_origen") or "",
                row.get("ciudad_destino") or "",
                precio_cliente,
                precio_tercero,
                rentabilidad,
                ganancia,
            ]
        )

        row_idx = ws.max_row
        for col_idx in range(1, 13):
            ws.cell(row=row_idx, column=col_idx).border = border

        for col_idx in (9, 10, 12):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.number_format = cop_format

        ws.cell(row=row_idx, column=11).number_format = '#,##0.##" %"'

    total_row = ws.max_row + 1
    ws.cell(row=total_row, column=8, value="TOTAL")
    ws.cell(row=total_row, column=9, value=total_precio_cliente)
    ws.cell(row=total_row, column=10, value=total_precio_tercero)
    ws.cell(row=total_row, column=12, value=total_ganancia)

    for col_idx in (8, 9, 10, 12):
        cell = ws.cell(row=total_row, column=col_idx)
        cell.fill = total_fill
        cell.font = total_font
        cell.border = border
        if col_idx in (9, 10, 12):
            cell.number_format = cop_format

    widths = {
        1: 14,
        2: 18,
        3: 24,
        4: 12,
        5: 12,
        6: 12,
        7: 16,
        8: 16,
        9: 18,
        10: 18,
        11: 14,
        12: 18,
    }
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _build_conciliacion_excel_legacy(
    conc: Conciliacion,
    items: list[ConciliacionItem],
    user_role: UserRole,
    tipo_vehiculo_by_placa: dict[str, str],
    avansat_prefetched: dict[str, dict] | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "resumen"

    header_fill = PatternFill(fill_type="solid", fgColor="E5E7EB")
    header_font = Font(bold=True, color="1F2937")
    section_fill = PatternFill(fill_type="solid", fgColor="FDE68A")
    section_font = Font(bold=True, color="111827")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    cop_format = '"$" #,##0'
    pct_format = '#,##0.##" %"'

    show_tarifa_tercero = user_role != UserRole.CLIENTE
    show_tarifa_cliente = user_role != UserRole.TERCERO
    show_cointra_financials = user_role == UserRole.COINTRA

    liquidacion_items: list[tuple[ConciliacionItem, dict]] = []
    additional_items: list[tuple[ConciliacionItem, dict | None]] = []
    for item in items:
        liq_meta = _extract_liquidacion_metadata(item)
        if liq_meta:
            liquidacion_items.append((item, liq_meta))
        else:
            service_code = str(item.servicio_codigo or "").strip().upper()
            if service_code == "DISPONIBILIDAD":
                continue
            additional_items.append((item, None))

    estado_display = _display_estado(conc)
    estado_label = estado_display.replace("_", " ")
    estado_styles = {
        "BORRADOR": ("ECFDF5", "065F46"),
        "EN_REVISION": ("FFFBEB", "92400E"),
        "APROBADA": ("F0FDFA", "115E59"),
        "ENVIADA_A_FACTURAR": ("F0F9FF", "075985"),
    }
    estado_fill_color, estado_font_color = estado_styles.get(estado_display, ("F7FEE7", "3F6212"))

    current_row = 1
    ws.cell(row=current_row, column=1, value=f"Conciliacion #{conc.id} - {conc.nombre}")
    ws.cell(row=current_row, column=1).font = Font(bold=True, size=12, color="111827")
    estado_cell = ws.cell(row=current_row, column=3, value=f"ESTADO: {estado_label}")
    estado_cell.fill = PatternFill(fill_type="solid", fgColor=estado_fill_color)
    estado_cell.font = Font(bold=True, color=estado_font_color)
    estado_cell.alignment = Alignment(horizontal="center", vertical="center")
    estado_cell.border = border
    current_row += 1
    ws.cell(row=current_row, column=1, value=f"Periodo: {conc.fecha_inicio} a {conc.fecha_fin}")
    ws.cell(row=current_row, column=1).font = Font(size=10, color="374151")
    current_row += 2

    # Bloque superior: liquidacion contrato fijo
    ws.cell(row=current_row, column=1, value="LIQUIDACION CONTRATO FIJO")
    ws.cell(row=current_row, column=1).font = Font(bold=True, size=11, color="111827")
    current_row += 1

    top_headers = ["Placa", "Tipo Vehiculo"]
    if show_tarifa_cliente:
        top_headers.append("Valor Cliente")
    if show_tarifa_tercero:
        top_headers.append("Valor Tercero")
    if show_cointra_financials:
        top_headers.extend(["Rentabilidad", "Ganancia Cointra"])

    top_col_idx = {header: idx for idx, header in enumerate(top_headers, start=1)}
    for idx, header in enumerate(top_headers, start=1):
        cell = ws.cell(row=current_row, column=idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border
    current_row += 1

    top_start_row = current_row
    liquidacion_items_sorted = sorted(
        liquidacion_items,
        key=lambda pair: (
            str(pair[0].placa or "").upper(),
            pair[0].fecha_servicio,
            pair[0].id,
        ),
    )

    for item, liq_meta in liquidacion_items_sorted:
        placa = (item.placa or "").upper()
        tipo_vehiculo = tipo_vehiculo_by_placa.get(placa, "")

        tarifa_tercero = _as_float(item.tarifa_tercero)
        tarifa_cliente = _as_float(item.tarifa_cliente)
        ganancia = tarifa_cliente - tarifa_tercero
        rentabilidad = _as_float(item.rentabilidad)

        ws.cell(row=current_row, column=top_col_idx["Placa"], value=placa)
        ws.cell(row=current_row, column=top_col_idx["Tipo Vehiculo"], value=tipo_vehiculo)
        if "Valor Cliente" in top_col_idx:
            c = ws.cell(row=current_row, column=top_col_idx["Valor Cliente"], value=tarifa_cliente)
            c.number_format = cop_format
        if "Valor Tercero" in top_col_idx:
            c = ws.cell(row=current_row, column=top_col_idx["Valor Tercero"], value=tarifa_tercero)
            c.number_format = cop_format
        if "Rentabilidad" in top_col_idx:
            c = ws.cell(row=current_row, column=top_col_idx["Rentabilidad"], value=rentabilidad)
            c.number_format = pct_format
        if "Ganancia Cointra" in top_col_idx:
            c = ws.cell(row=current_row, column=top_col_idx["Ganancia Cointra"], value=ganancia)
            c.number_format = cop_format

        for col_idx in range(1, len(top_headers) + 1):
            ws.cell(row=current_row, column=col_idx).border = border
        current_row += 1

    if current_row == top_start_row:
        ws.cell(row=current_row, column=1, value="(sin registros de contrato fijo)")
        ws.cell(row=current_row, column=1).font = Font(italic=True, color="6B7280")
        current_row += 1

    top_total_tercero = sum(_as_float(item.tarifa_tercero) for item, _ in liquidacion_items_sorted)
    top_total_cliente = sum(_as_float(item.tarifa_cliente) for item, _ in liquidacion_items_sorted)
    top_ganancia = top_total_cliente - top_total_tercero
    top_rentabilidad = (top_ganancia / top_total_cliente * 100) if top_total_cliente > 0 else 0.0

    current_row += 1
    ws.cell(row=current_row, column=2, value="TOTAL CONTRATO FIJO")
    if "Valor Cliente" in top_col_idx:
        c = ws.cell(row=current_row, column=top_col_idx["Valor Cliente"], value=top_total_cliente)
        c.number_format = cop_format
    if "Valor Tercero" in top_col_idx:
        c = ws.cell(row=current_row, column=top_col_idx["Valor Tercero"], value=top_total_tercero)
        c.number_format = cop_format
    if "Rentabilidad" in top_col_idx:
        c = ws.cell(row=current_row, column=top_col_idx["Rentabilidad"], value=top_rentabilidad)
        c.number_format = pct_format
    if "Ganancia Cointra" in top_col_idx:
        c = ws.cell(row=current_row, column=top_col_idx["Ganancia Cointra"], value=top_ganancia)
        c.number_format = cop_format

    for col_idx in range(1, len(top_headers) + 1):
        cell = ws.cell(row=current_row, column=col_idx)
        cell.fill = section_fill
        cell.font = section_font
        cell.border = border

    # Bloque inferior: adicionales/otros servicios
    current_row += 2
    ws.cell(row=current_row, column=1, value="SERVICIOS")
    ws.cell(row=current_row, column=1).font = Font(bold=True, size=11, color="111827")
    current_row += 1

    bottom_headers = ["Placa", "Tipo Vehiculo", "Fecha", "Titulo Servicio", "Tipo Servicio"]
    if show_tarifa_cliente:
        bottom_headers.append("Valor Cliente")
    if show_tarifa_tercero:
        bottom_headers.append("Valor Tercero")
    if show_cointra_financials:
        bottom_headers.extend(["Rentabilidad", "Ganancia Cointra"])
    bottom_headers.append("Observaciones")
    bottom_col_idx = {header: idx for idx, header in enumerate(bottom_headers, start=1)}

    for idx, header in enumerate(bottom_headers, start=1):
        cell = ws.cell(row=current_row, column=idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border
    current_row += 1

    additional_sorted = sorted(
        additional_items,
        key=lambda pair: (
            str(pair[0].placa or "").upper(),
            pair[0].fecha_servicio,
            pair[0].id,
        ),
    )

    group_placa = None
    group_total_tercero = 0.0
    group_total_cliente = 0.0
    all_total_tercero = 0.0
    all_total_cliente = 0.0

    def _write_group_total(row_num: int, placa: str, total_tercero: float, total_cliente: float) -> int:
        if not placa:
            return row_num
        ganancia = total_cliente - total_tercero
        rentabilidad = (ganancia / total_cliente * 100) if total_cliente > 0 else 0.0
        ws.cell(row=row_num, column=2, value=f"TOTAL PLACA {placa}")
        if "Valor Cliente" in bottom_col_idx:
            c = ws.cell(row=row_num, column=bottom_col_idx["Valor Cliente"], value=total_cliente)
            c.number_format = cop_format
        if "Valor Tercero" in bottom_col_idx:
            c = ws.cell(row=row_num, column=bottom_col_idx["Valor Tercero"], value=total_tercero)
            c.number_format = cop_format
        if "Rentabilidad" in bottom_col_idx:
            c = ws.cell(row=row_num, column=bottom_col_idx["Rentabilidad"], value=rentabilidad)
            c.number_format = pct_format
        if "Ganancia Cointra" in bottom_col_idx:
            c = ws.cell(row=row_num, column=bottom_col_idx["Ganancia Cointra"], value=ganancia)
            c.number_format = cop_format

        for col_idx in range(1, len(bottom_headers) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = section_font
            cell.border = border
        return row_num + 1

    for item, _ in additional_sorted:
        placa = (item.placa or "").upper() or "SIN_PLACA"
        if group_placa is None:
            group_placa = placa
        elif placa != group_placa:
            current_row = _write_group_total(current_row, group_placa, group_total_tercero, group_total_cliente)
            group_placa = placa
            group_total_tercero = 0.0
            group_total_cliente = 0.0

        tipo_vehiculo = tipo_vehiculo_by_placa.get(placa, "")
        titulo_servicio = item.viaje.titulo if item.viaje else ""
        tipo_servicio = (
            (item.servicio_nombre or "").strip()
            or (item.servicio_codigo or "").strip()
            or str(getattr(item.tipo, "value", item.tipo))
        )
        tarifa_tercero = _as_float(item.tarifa_tercero)
        tarifa_cliente = _as_float(item.tarifa_cliente)
        ganancia = tarifa_cliente - tarifa_tercero
        rentabilidad = _as_float(item.rentabilidad)

        group_total_tercero += tarifa_tercero
        group_total_cliente += tarifa_cliente
        all_total_tercero += tarifa_tercero
        all_total_cliente += tarifa_cliente

        ws.cell(row=current_row, column=bottom_col_idx["Placa"], value=placa)
        ws.cell(row=current_row, column=bottom_col_idx["Tipo Vehiculo"], value=tipo_vehiculo)
        ws.cell(row=current_row, column=bottom_col_idx["Fecha"], value=str(item.fecha_servicio))
        ws.cell(row=current_row, column=bottom_col_idx["Titulo Servicio"], value=titulo_servicio)
        ws.cell(row=current_row, column=bottom_col_idx["Tipo Servicio"], value=tipo_servicio)
        if "Valor Cliente" in bottom_col_idx:
            c = ws.cell(row=current_row, column=bottom_col_idx["Valor Cliente"], value=tarifa_cliente)
            c.number_format = cop_format
        if "Valor Tercero" in bottom_col_idx:
            c = ws.cell(row=current_row, column=bottom_col_idx["Valor Tercero"], value=tarifa_tercero)
            c.number_format = cop_format
        if "Rentabilidad" in bottom_col_idx:
            c = ws.cell(row=current_row, column=bottom_col_idx["Rentabilidad"], value=rentabilidad)
            c.number_format = pct_format
        if "Ganancia Cointra" in bottom_col_idx:
            c = ws.cell(row=current_row, column=bottom_col_idx["Ganancia Cointra"], value=ganancia)
            c.number_format = cop_format
        ws.cell(row=current_row, column=bottom_col_idx["Observaciones"], value=item.descripcion or "")

        for col_idx in range(1, len(bottom_headers) + 1):
            ws.cell(row=current_row, column=col_idx).border = border
        current_row += 1

    if not additional_sorted:
        ws.cell(row=current_row, column=1, value="(sin servicios adicionales)")
        ws.cell(row=current_row, column=1).font = Font(italic=True, color="6B7280")
        current_row += 1
    else:
        current_row = _write_group_total(current_row, group_placa or "", group_total_tercero, group_total_cliente)

        all_ganancia = all_total_cliente - all_total_tercero
        all_rentabilidad = (all_ganancia / all_total_cliente * 100) if all_total_cliente > 0 else 0.0
        ws.cell(row=current_row, column=2, value="TOTAL SERVICIOS")
        if "Valor Cliente" in bottom_col_idx:
            c = ws.cell(row=current_row, column=bottom_col_idx["Valor Cliente"], value=all_total_cliente)
            c.number_format = cop_format
        if "Valor Tercero" in bottom_col_idx:
            c = ws.cell(row=current_row, column=bottom_col_idx["Valor Tercero"], value=all_total_tercero)
            c.number_format = cop_format
        if "Rentabilidad" in bottom_col_idx:
            c = ws.cell(row=current_row, column=bottom_col_idx["Rentabilidad"], value=all_rentabilidad)
            c.number_format = pct_format
        if "Ganancia Cointra" in bottom_col_idx:
            c = ws.cell(row=current_row, column=bottom_col_idx["Ganancia Cointra"], value=all_ganancia)
            c.number_format = cop_format
        for col_idx in range(1, len(bottom_headers) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = section_fill
            cell.font = section_font
            cell.border = border
        current_row += 2

    # Hoja 2: servicios (manifiestos del contexto conciliacion + datos Avansat)
    manifest_financials_servicios: dict[str, dict[str, float]] = {}
    servicios_totals_by_placa: dict[str, dict[str, float]] = {}
    for additional_item, _ in additional_items:
        manifest_key = _normalize_manifiesto_for_lookup(additional_item.manifiesto_numero)
        if manifest_key:
            bucket = manifest_financials_servicios.setdefault(
                manifest_key,
                {"valor_tercero": 0.0, "valor_cliente": 0.0},
            )
            bucket["valor_tercero"] += _as_float(additional_item.tarifa_tercero)
            bucket["valor_cliente"] += _as_float(additional_item.tarifa_cliente)

        placa_key = str(additional_item.placa or "").strip().upper()
        if placa_key:
            plate_bucket = servicios_totals_by_placa.setdefault(
                placa_key,
                {"valor_tercero": 0.0, "valor_cliente": 0.0},
            )
            plate_bucket["valor_tercero"] += _as_float(additional_item.tarifa_tercero)
            plate_bucket["valor_cliente"] += _as_float(additional_item.tarifa_cliente)

    ws_servicios = wb.create_sheet("Servicios")
    servicios_headers = [
        "Manifiesto",
        "Fecha Emision",
        "Placa Vehiculo",
        "Trayler",
        "Remesa",
        "Producto",
        "Ciudad Origen",
        "Ciudad Destino",
    ]
    if show_tarifa_cliente:
        servicios_headers.append("Valor Cliente")
    if show_tarifa_tercero:
        servicios_headers.append("Valor Tercero")
    if show_cointra_financials:
        servicios_headers.extend(["Rentabilidad", "Ganancia Cointra"])

    servicios_col_idx = {header: idx for idx, header in enumerate(servicios_headers, start=1)}
    ws_servicios.append(servicios_headers)
    for idx, cell in enumerate(ws_servicios[1], start=1):
        cell.fill = header_fill
        if servicios_headers[idx - 1] in {"Valor Tercero", "Valor Cliente", "Rentabilidad", "Ganancia Cointra"}:
            cell.fill = section_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    def _write_servicios_block_header(row_idx: int) -> None:
        for idx, header in enumerate(servicios_headers, start=1):
            cell = ws_servicios.cell(row=row_idx, column=idx, value=header)
            cell.fill = header_fill
            if header in {"Valor Tercero", "Valor Cliente", "Rentabilidad", "Ganancia Cointra"}:
                cell.fill = section_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = border

    servicios_entries: list[dict[str, object]] = []
    manifests_sorted = sorted(manifest_financials_servicios.keys())
    avansat_servicios_prefetched = avansat_prefetched or {}

    for manifest_key in manifests_sorted:
        manifiesto = _normalize_manifiesto_for_lookup(manifest_key)
        avansat = avansat_servicios_prefetched.get(manifiesto) or {}
        remesas = avansat.get("remesas") if isinstance(avansat.get("remesas"), list) else []
        remesas_rows = [r for r in remesas if isinstance(r, dict)]
        if not remesas_rows:
            remesas_rows = [{"remesa": avansat.get("remesa") or "", "producto": avansat.get("producto") or ""}]

        financials = manifest_financials_servicios.get(manifiesto, {"valor_tercero": 0.0, "valor_cliente": 0.0})
        valor_tercero = _as_float(financials.get("valor_tercero"))
        valor_cliente = _as_float(financials.get("valor_cliente"))
        ganancia = valor_cliente - valor_tercero
        rentabilidad = (ganancia / valor_cliente * 100) if valor_cliente > 0 else 0.0
        placa = str(avansat.get("placa_vehiculo") or "").strip().upper() or "SIN_PLACA"

        servicios_entries.append(
            {
                "placa": placa,
                "manifiesto": manifiesto,
                "fecha_emision": avansat.get("fecha_emision") or "",
                "trayler": avansat.get("trayler") or "",
                "ciudad_origen": avansat.get("ciudad_origen") or "",
                "ciudad_destino": avansat.get("ciudad_destino") or "",
                "remesas": remesas_rows,
                "valor_tercero": valor_tercero,
                "valor_cliente": valor_cliente,
                "rentabilidad": rentabilidad,
                "ganancia": ganancia,
            }
        )

    servicios_entries.sort(key=lambda row: (str(row["placa"]), str(row["manifiesto"])))

    if not servicios_entries:
        ws_servicios.append(["Sin manifiestos asociados a conciliacion"])
        ws_servicios.cell(row=2, column=1).font = Font(italic=True, color="6B7280")
    else:
        current_row = 2
        total_general_tercero = 0.0
        total_general_cliente = 0.0
        placas = sorted({str(row["placa"]) for row in servicios_entries})

        for plate_index, placa in enumerate(placas):
            plate_rows = [row for row in servicios_entries if str(row["placa"]) == placa]
            total_placa_tercero = 0.0
            total_placa_cliente = 0.0

            servicios_totals = servicios_totals_by_placa.get(placa)
            if servicios_totals:
                total_placa_tercero = _as_float(servicios_totals.get("valor_tercero"))
                total_placa_cliente = _as_float(servicios_totals.get("valor_cliente"))
            else:
                # Fallback: si no hay registro de servicios para la placa, conserva suma por manifiestos.
                total_placa_tercero = sum(_as_float(row.get("valor_tercero")) for row in plate_rows)
                total_placa_cliente = sum(_as_float(row.get("valor_cliente")) for row in plate_rows)

            manifest_count = len(plate_rows)
            per_manifest_tercero = _split_total_evenly(total_placa_tercero, manifest_count)
            per_manifest_cliente = _split_total_evenly(total_placa_cliente, manifest_count)

            if plate_index > 0:
                current_row += 1
                _write_servicios_block_header(current_row)
                current_row += 1

            for entry_idx, entry in enumerate(plate_rows):
                remesas_rows = entry["remesas"] if isinstance(entry["remesas"], list) else []
                manifest_valor_tercero = per_manifest_tercero[entry_idx] if entry_idx < len(per_manifest_tercero) else 0.0
                manifest_valor_cliente = per_manifest_cliente[entry_idx] if entry_idx < len(per_manifest_cliente) else 0.0
                manifest_ganancia = manifest_valor_cliente - manifest_valor_tercero
                manifest_rentabilidad = (
                    (manifest_ganancia / manifest_valor_cliente * 100) if manifest_valor_cliente > 0 else 0.0
                )
                first_row_for_manifest = True
                for remesa_row in remesas_rows:
                    row_values: dict[str, object] = {
                        "Manifiesto": entry["manifiesto"],
                        "Fecha Emision": entry["fecha_emision"],
                        "Placa Vehiculo": entry["placa"],
                        "Trayler": entry["trayler"],
                        "Remesa": str((remesa_row or {}).get("remesa") or "").strip(),
                        "Producto": str((remesa_row or {}).get("producto") or "").strip(),
                        "Ciudad Origen": entry["ciudad_origen"],
                        "Ciudad Destino": entry["ciudad_destino"],
                    }
                    if show_tarifa_tercero:
                        row_values["Valor Tercero"] = manifest_valor_tercero if first_row_for_manifest else None
                    if show_tarifa_cliente:
                        row_values["Valor Cliente"] = manifest_valor_cliente if first_row_for_manifest else None
                    if show_cointra_financials:
                        row_values["Rentabilidad"] = manifest_rentabilidad if first_row_for_manifest else None
                        row_values["Ganancia Cointra"] = manifest_ganancia if first_row_for_manifest else None

                    for header, col_idx in servicios_col_idx.items():
                        cell = ws_servicios.cell(row=current_row, column=col_idx, value=row_values.get(header, ""))
                        cell.border = border
                        if header in {"Valor Tercero", "Valor Cliente", "Ganancia Cointra"} and cell.value is not None:
                            cell.number_format = cop_format
                        if header == "Rentabilidad" and cell.value is not None:
                            cell.number_format = pct_format
                    current_row += 1
                    first_row_for_manifest = False

            total_general_tercero += total_placa_tercero
            total_general_cliente += total_placa_cliente

            ws_servicios.cell(row=current_row, column=servicios_col_idx["Ciudad Destino"], value=f"TOTAL PLACA {placa}")
            total_placa_cells_to_fill = [servicios_col_idx["Ciudad Destino"]]
            if "Valor Tercero" in servicios_col_idx:
                c = ws_servicios.cell(row=current_row, column=servicios_col_idx["Valor Tercero"], value=total_placa_tercero)
                c.number_format = cop_format
                total_placa_cells_to_fill.append(servicios_col_idx["Valor Tercero"])
            if "Valor Cliente" in servicios_col_idx:
                c = ws_servicios.cell(row=current_row, column=servicios_col_idx["Valor Cliente"], value=total_placa_cliente)
                c.number_format = cop_format
                total_placa_cells_to_fill.append(servicios_col_idx["Valor Cliente"])
            if "Rentabilidad" in servicios_col_idx:
                placa_ganancia = total_placa_cliente - total_placa_tercero
                placa_rentabilidad = (placa_ganancia / total_placa_cliente * 100) if total_placa_cliente > 0 else 0.0
                c = ws_servicios.cell(row=current_row, column=servicios_col_idx["Rentabilidad"], value=placa_rentabilidad)
                c.number_format = pct_format
                total_placa_cells_to_fill.append(servicios_col_idx["Rentabilidad"])
            if "Ganancia Cointra" in servicios_col_idx:
                c = ws_servicios.cell(
                    row=current_row,
                    column=servicios_col_idx["Ganancia Cointra"],
                    value=total_placa_cliente - total_placa_tercero,
                )
                c.number_format = cop_format
                total_placa_cells_to_fill.append(servicios_col_idx["Ganancia Cointra"])

            for col_idx in range(1, len(servicios_headers) + 1):
                cell = ws_servicios.cell(row=current_row, column=col_idx)
                cell.font = section_font
                cell.border = border
            for col_idx in total_placa_cells_to_fill:
                ws_servicios.cell(row=current_row, column=col_idx).fill = section_fill
            current_row += 1

        current_row += 2
        ws_servicios.cell(row=current_row, column=servicios_col_idx["Ciudad Destino"], value="TOTAL GENERAL")
        total_general_cells_to_fill = [servicios_col_idx["Ciudad Destino"]]
        if "Valor Tercero" in servicios_col_idx:
            c = ws_servicios.cell(row=current_row, column=servicios_col_idx["Valor Tercero"], value=total_general_tercero)
            c.number_format = cop_format
            total_general_cells_to_fill.append(servicios_col_idx["Valor Tercero"])
        if "Valor Cliente" in servicios_col_idx:
            c = ws_servicios.cell(row=current_row, column=servicios_col_idx["Valor Cliente"], value=total_general_cliente)
            c.number_format = cop_format
            total_general_cells_to_fill.append(servicios_col_idx["Valor Cliente"])
        if "Rentabilidad" in servicios_col_idx:
            total_ganancia = total_general_cliente - total_general_tercero
            total_rentabilidad = (total_ganancia / total_general_cliente * 100) if total_general_cliente > 0 else 0.0
            c = ws_servicios.cell(row=current_row, column=servicios_col_idx["Rentabilidad"], value=total_rentabilidad)
            c.number_format = pct_format
            total_general_cells_to_fill.append(servicios_col_idx["Rentabilidad"])
        if "Ganancia Cointra" in servicios_col_idx:
            c = ws_servicios.cell(
                row=current_row,
                column=servicios_col_idx["Ganancia Cointra"],
                value=total_general_cliente - total_general_tercero,
            )
            c.number_format = cop_format
            total_general_cells_to_fill.append(servicios_col_idx["Ganancia Cointra"])

        for col_idx in range(1, len(servicios_headers) + 1):
            cell = ws_servicios.cell(row=current_row, column=col_idx)
            cell.font = section_font
            cell.border = border
        for col_idx in total_general_cells_to_fill:
            ws_servicios.cell(row=current_row, column=col_idx).fill = section_fill

    servicios_widths = {
        "Manifiesto": 20,
        "Fecha Emision": 18,
        "Placa Vehiculo": 20,
        "Trayler": 18,
        "Remesa": 20,
        "Producto": 40,
        "Ciudad Origen": 24,
        "Ciudad Destino": 24,
        "Valor Tercero": 20,
        "Valor Cliente": 20,
        "Rentabilidad": 16,
        "Ganancia Cointra": 22,
    }
    for header, idx in servicios_col_idx.items():
        ws_servicios.column_dimensions[get_column_letter(idx)].width = servicios_widths.get(header, 18)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
