from typing import NamedTuple
from openpyxl.styles import Alignment, Border, Font, PatternFill
from app.models.conciliacion_item import ConciliacionItem
from .conciliaciones_helpers import (
    _as_float,
    _item_servicio_codigo,
    _normalize_manifiesto_for_lookup,
)


class _ExcelContext(NamedTuple):
    conc: object                        # instancia Conciliacion ORM
    avansat_prefetched: dict | None
    show_tarifa_tercero: bool
    show_tarifa_cliente: bool
    show_cointra_financials: bool
    border: Border
    cop_format: str
    pct_format: str
    header_fill: PatternFill
    header_font: Font
    section_fill: PatternFill
    section_font: Font
    center: Alignment
    tipo_vehiculo_by_placa: dict[str, str]
    estado_label: str
    estado_fill_color: str
    estado_font_color: str


def _xl_write_report_header(ws, title: str, ctx: _ExcelContext) -> int:
    row_num = 1
    ws.cell(row=row_num, column=1, value=f"Conciliacion #{ctx.conc.id} - {ctx.conc.nombre}")
    ws.cell(row=row_num, column=1).font = Font(bold=True, size=12, color="111827")
    ws.cell(row=row_num, column=1).alignment = Alignment(horizontal="left", vertical="center")
    title_cell = ws.cell(row=row_num, column=5, value=title)
    title_cell.font = Font(bold=True, size=12, color="111827")
    estado_cell = ws.cell(row=row_num, column=7, value=f"ESTADO: {ctx.estado_label}")
    estado_cell.fill = PatternFill(fill_type="solid", fgColor=ctx.estado_fill_color)
    estado_cell.font = Font(bold=True, color=ctx.estado_font_color)
    estado_cell.alignment = ctx.center
    estado_cell.border = ctx.border
    row_num += 1
    ws.cell(row=row_num, column=1, value=f"Periodo: {ctx.conc.fecha_inicio} a {ctx.conc.fecha_fin}")
    ws.cell(row=row_num, column=1).font = Font(size=10, color="374151")
    return row_num + 2


def _xl_write_headers(ws, row_num: int, headers: list[str], ctx: _ExcelContext) -> dict[str, int]:
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_num, column=idx, value=header)
        cell.fill = ctx.header_fill
        if header in {"Valor Cliente", "Valor Tercero", "Rentabilidad", "Ganancia Cointra"}:
            cell.fill = ctx.section_fill
        cell.font = ctx.header_font
        cell.alignment = ctx.center
        cell.border = ctx.border
    return {header: idx for idx, header in enumerate(headers, start=1)}


def _xl_style_row(
    ws,
    row_num: int,
    columns_count: int,
    ctx: _ExcelContext,
    fill: PatternFill | None = None,
    font: Font | None = None,
) -> None:
    for col_idx in range(1, columns_count + 1):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.border = ctx.border
        if fill is not None:
            cell.fill = fill
        if font is not None:
            cell.font = font


def _xl_write_financials(
    ws,
    row_num: int,
    col_idx: dict[str, int],
    tarifa_cliente: float,
    tarifa_tercero: float,
    ctx: _ExcelContext,
) -> None:
    ganancia = tarifa_cliente - tarifa_tercero
    rentabilidad = (ganancia / tarifa_cliente * 100) if tarifa_cliente > 0 else 0.0
    if "Valor Cliente" in col_idx:
        c = ws.cell(row=row_num, column=col_idx["Valor Cliente"], value=tarifa_cliente)
        c.number_format = ctx.cop_format
    if "Valor Tercero" in col_idx:
        c = ws.cell(row=row_num, column=col_idx["Valor Tercero"], value=tarifa_tercero)
        c.number_format = ctx.cop_format
    if "Rentabilidad" in col_idx:
        c = ws.cell(row=row_num, column=col_idx["Rentabilidad"], value=rentabilidad)
        c.number_format = ctx.pct_format
    if "Ganancia Cointra" in col_idx:
        c = ws.cell(row=row_num, column=col_idx["Ganancia Cointra"], value=ganancia)
        c.number_format = ctx.cop_format


def _xl_accumulate_totals(source_items: list[ConciliacionItem]) -> dict[str, dict[str, float]]:
    totals: dict[str, dict[str, float]] = {}
    for source_item in source_items:
        placa = str(source_item.placa or "").strip().upper() or "SIN_PLACA"
        bucket = totals.setdefault(placa, {"cliente": 0.0, "tercero": 0.0})
        bucket["cliente"] += _as_float(source_item.tarifa_cliente)
        bucket["tercero"] += _as_float(source_item.tarifa_tercero)
    return totals


def _xl_manifest_context(item: ConciliacionItem, ctx: _ExcelContext) -> dict[str, object]:
    manifiesto = _normalize_manifiesto_for_lookup(item.manifiesto_numero)
    avansat = (ctx.avansat_prefetched or {}).get(manifiesto) or {}
    has_manifiesto = bool(manifiesto)
    remesas = avansat.get("remesas") if isinstance(avansat.get("remesas"), list) else []
    remesas_rows = [row for row in remesas if isinstance(row, dict)]
    if not remesas_rows:
        remesas_rows = [
            {
                "remesa": item.remesa or avansat.get("remesa") or "",
                "producto": avansat.get("producto") or "",
            }
        ]
    placa = str(item.placa or "").strip().upper() or "SIN_PLACA"
    return {
        "manifiesto": manifiesto,
        "has_manifiesto": has_manifiesto,
        "fecha_emision": avansat.get("fecha_emision") or str(item.fecha_servicio or ""),
        "placa": placa,
        "tipo_vehiculo": ctx.tipo_vehiculo_by_placa.get(placa, ""),
        "trayler": avansat.get("trayler") or "",
        "ciudad_origen": avansat.get("ciudad_origen") or item.origen or "",
        "ciudad_destino": avansat.get("ciudad_destino") or item.destino or "",
        "fecha_cumplida": avansat.get("fecha_cumplida") or "",
        "remesas": remesas_rows,
    }


def _xl_item_service_code_upper(item: ConciliacionItem) -> str:
    codigo_viaje = _item_servicio_codigo(item)
    if codigo_viaje:
        return codigo_viaje
    return str(item.servicio_codigo or "").strip().upper()


def _xl_item_service_label(item: ConciliacionItem) -> str:
    return (
        (item.servicio_nombre or "").strip()
        or _xl_item_service_code_upper(item)
        or str(getattr(item.tipo, "value", item.tipo))
    )


def _xl_write_adicionales_unified_section(
    ws,
    start_row: int,
    title: str,
    source_items: list[ConciliacionItem],
    ctx: _ExcelContext,
) -> tuple[int, dict[str, int], float, float, float, float]:
    ws.cell(row=start_row, column=1, value=title)
    ws.cell(row=start_row, column=1).font = Font(bold=True, size=11, color="111827")
    headers = [
        "Manifiesto",
        "Fecha Emision",
        "Fecha Cumplida",
        "Placa Vehiculo",
        "Tipo Vehiculo",
        "Trayler",
        "Remesa",
        "Producto",
        "Ciudad Origen",
        "Ciudad Destino",
    ]
    if ctx.show_tarifa_cliente:
        headers.append("Valor Cliente")
    if ctx.show_tarifa_tercero:
        headers.append("Valor Tercero")
    if ctx.show_cointra_financials:
        headers.extend(["Rentabilidad", "Ganancia Cointra"])

    col_idx = _xl_write_headers(ws, start_row + 1, headers, ctx)
    row_num = start_row + 2
    total_viajes_cliente = 0.0
    total_viajes_tercero = 0.0
    total_servicios_cliente = 0.0
    total_servicios_tercero = 0.0

    sorted_items = sorted(
        source_items,
        key=lambda current: (
            str(current.placa or "").strip().upper(),
            current.fecha_servicio,
            current.id,
        ),
    )

    if not sorted_items:
        ws.cell(row=row_num, column=1, value="(sin adicionales)")
        ws.cell(row=row_num, column=1).font = Font(italic=True, color="6B7280")
        return row_num + 1, col_idx, total_viajes_cliente, total_viajes_tercero, total_servicios_cliente, total_servicios_tercero

    grouped_items: dict[str, list[ConciliacionItem]] = {}
    for item in sorted_items:
        placa = str(item.placa or "").strip().upper() or "SIN_PLACA"
        grouped_items.setdefault(placa, []).append(item)

    ordered_placas = sorted(grouped_items.keys())
    for placa_idx, placa in enumerate(ordered_placas):
        if placa_idx > 0:
            _xl_write_headers(ws, row_num, headers, ctx)
            row_num += 1
        placa_total_cliente = 0.0
        placa_total_tercero = 0.0

        for item in grouped_items[placa]:
            manifest_data = _xl_manifest_context(item, ctx)
            has_manifiesto = bool(manifest_data["has_manifiesto"])
            manifiesto_str = str(manifest_data["manifiesto"] or "")
            service_label = _xl_item_service_label(item) if not has_manifiesto else ""

            tarifa_cliente = _as_float(item.tarifa_cliente)
            tarifa_tercero = _as_float(item.tarifa_tercero)
            placa_total_cliente += tarifa_cliente
            placa_total_tercero += tarifa_tercero

            if has_manifiesto:
                total_viajes_cliente += tarifa_cliente
                total_viajes_tercero += tarifa_tercero
            else:
                total_servicios_cliente += tarifa_cliente
                total_servicios_tercero += tarifa_tercero

            first_row = True
            for remesa_row in manifest_data["remesas"]:
                values = {
                    "Manifiesto": manifiesto_str if first_row else "",
                    "Fecha Emision": manifest_data["fecha_emision"],
                    "Fecha Cumplida": manifest_data["fecha_cumplida"],
                    "Placa Vehiculo": placa,
                    "Tipo Vehiculo": manifest_data["tipo_vehiculo"],
                    "Trayler": manifest_data["trayler"],
                    "Remesa": str((remesa_row or {}).get("remesa") or "").strip(),
                    "Producto": service_label if (first_row and not has_manifiesto) else str((remesa_row or {}).get("producto") or "").strip(),
                    "Ciudad Origen": manifest_data["ciudad_origen"],
                    "Ciudad Destino": manifest_data["ciudad_destino"],
                }
                for header, column in col_idx.items():
                    if header in values:
                        ws.cell(row=row_num, column=column, value=values[header])
                if first_row:
                    _xl_write_financials(ws, row_num, col_idx, tarifa_cliente, tarifa_tercero, ctx)
                _xl_style_row(ws, row_num, len(headers), ctx)
                row_num += 1
                first_row = False

        ws.cell(row=row_num, column=col_idx["Ciudad Destino"], value=f"TOTAL PLACA {placa}")
        _xl_write_financials(ws, row_num, col_idx, placa_total_cliente, placa_total_tercero, ctx)
        _xl_style_row(ws, row_num, len(headers), ctx, fill=ctx.section_fill, font=ctx.section_font)
        row_num += 2

    return row_num, col_idx, total_viajes_cliente, total_viajes_tercero, total_servicios_cliente, total_servicios_tercero


def _xl_write_transport_section(
    ws,
    start_row: int,
    title: str,
    source_items: list[ConciliacionItem],
    ctx: _ExcelContext,
    write_plate_totals: bool = False,
    totals_label_prefix: str = "TOTAL PLACA",
) -> tuple[int, dict[str, int], float, float, dict[str, dict[str, float]]]:
    ws.cell(row=start_row, column=1, value=title)
    ws.cell(row=start_row, column=1).font = Font(bold=True, size=11, color="111827")
    headers = [
        "Manifiesto",
        "Fecha Emision",
        "Fecha Cumplida",
        "Placa Vehiculo",
        "Tipo Vehiculo",
        "Trayler",
        "Remesa",
        "Producto",
        "Ciudad Origen",
        "Ciudad Destino",
    ]
    if ctx.show_tarifa_cliente:
        headers.append("Valor Cliente")
    if ctx.show_tarifa_tercero:
        headers.append("Valor Tercero")
    if ctx.show_cointra_financials:
        headers.extend(["Rentabilidad", "Ganancia Cointra"])

    col_idx = _xl_write_headers(ws, start_row + 1, headers, ctx)
    row_num = start_row + 2
    total_cliente = 0.0
    total_tercero = 0.0
    totals_by_placa = _xl_accumulate_totals(source_items)
    sorted_items = sorted(
        source_items,
        key=lambda current: (
            str(current.placa or "").strip().upper(),
            current.fecha_servicio,
            current.id,
        ),
    )

    if not sorted_items:
        ws.cell(row=row_num, column=1, value="(sin registros)")
        ws.cell(row=row_num, column=1).font = Font(italic=True, color="6B7280")
        return row_num + 1, col_idx, total_cliente, total_tercero, totals_by_placa

    if write_plate_totals:
        grouped_items: dict[str, list[ConciliacionItem]] = {}
        for item in sorted_items:
            placa = str(item.placa or "").strip().upper() or "SIN_PLACA"
            grouped_items.setdefault(placa, []).append(item)

        ordered_placas = sorted(grouped_items.keys())
        for placa_idx, placa in enumerate(ordered_placas):
            if placa_idx > 0:
                _xl_write_headers(ws, row_num, headers, ctx)
                row_num += 1
            for item in grouped_items[placa]:
                manifest_data = _xl_manifest_context(item, ctx)
                tarifa_cliente = _as_float(item.tarifa_cliente)
                tarifa_tercero = _as_float(item.tarifa_tercero)
                total_cliente += tarifa_cliente
                total_tercero += tarifa_tercero

                first_row = True
                for remesa_row in manifest_data["remesas"]:
                    values = {
                        "Manifiesto": manifest_data["manifiesto"],
                        "Fecha Emision": manifest_data["fecha_emision"],
                        "Fecha Cumplida": manifest_data["fecha_cumplida"],
                        "Placa Vehiculo": placa,
                        "Tipo Vehiculo": manifest_data["tipo_vehiculo"],
                        "Trayler": manifest_data["trayler"],
                        "Remesa": str((remesa_row or {}).get("remesa") or "").strip(),
                        "Producto": str((remesa_row or {}).get("producto") or "").strip(),
                        "Ciudad Origen": manifest_data["ciudad_origen"],
                        "Ciudad Destino": manifest_data["ciudad_destino"],
                    }
                    for header, column in col_idx.items():
                        if header in values:
                            ws.cell(row=row_num, column=column, value=values[header])
                    if first_row:
                        _xl_write_financials(ws, row_num, col_idx, tarifa_cliente, tarifa_tercero, ctx)
                    _xl_style_row(ws, row_num, len(headers), ctx)
                    row_num += 1
                    first_row = False

            placa_totals = totals_by_placa.get(placa, {"cliente": 0.0, "tercero": 0.0})
            ws.cell(row=row_num, column=col_idx["Ciudad Destino"], value=f"{totals_label_prefix} {placa}")
            _xl_write_financials(
                ws,
                row_num,
                col_idx,
                placa_totals["cliente"],
                placa_totals["tercero"],
                ctx,
            )
            _xl_style_row(ws, row_num, len(headers), ctx, fill=ctx.section_fill, font=ctx.section_font)
            row_num += 2

        return row_num, col_idx, total_cliente, total_tercero, totals_by_placa

    current_placa = None
    for item in sorted_items:
        manifest_data = _xl_manifest_context(item, ctx)
        placa = str(manifest_data["placa"])
        if current_placa is not None and placa != current_placa:
            row_num += 1  # fila en blanco de separación
            _xl_write_headers(ws, row_num, headers, ctx)
            row_num += 1
        current_placa = placa

        tarifa_cliente = _as_float(item.tarifa_cliente)
        tarifa_tercero = _as_float(item.tarifa_tercero)
        total_cliente += tarifa_cliente
        total_tercero += tarifa_tercero

        first_row = True
        for remesa_row in manifest_data["remesas"]:
            values = {
                "Manifiesto": manifest_data["manifiesto"],
                "Fecha Emision": manifest_data["fecha_emision"],
                "Fecha Cumplida": manifest_data["fecha_cumplida"],
                "Placa Vehiculo": placa,
                "Tipo Vehiculo": manifest_data["tipo_vehiculo"],
                "Trayler": manifest_data["trayler"],
                "Remesa": str((remesa_row or {}).get("remesa") or "").strip(),
                "Producto": str((remesa_row or {}).get("producto") or "").strip(),
                "Ciudad Origen": manifest_data["ciudad_origen"],
                "Ciudad Destino": manifest_data["ciudad_destino"],
            }
            for header, column in col_idx.items():
                if header in values:
                    ws.cell(row=row_num, column=column, value=values[header])
            if first_row:
                _xl_write_financials(ws, row_num, col_idx, tarifa_cliente, tarifa_tercero, ctx)
            _xl_style_row(ws, row_num, len(headers), ctx)
            row_num += 1
            first_row = False

    return row_num, col_idx, total_cliente, total_tercero, totals_by_placa


def _xl_write_additional_services_section(
    ws,
    start_row: int,
    title: str,
    source_items: list[ConciliacionItem],
    ctx: _ExcelContext,
    write_plate_totals: bool = False,
    totals_label_prefix: str = "TOTAL PLACA",
) -> tuple[int, dict[str, int], float, float]:
    ws.cell(row=start_row, column=1, value=title)
    ws.cell(row=start_row, column=1).font = Font(bold=True, size=11, color="111827")
    headers = ["Placa", "Tipo Vehiculo", "Fecha", "Titulo Servicio", "Tipo Servicio"]
    if ctx.show_tarifa_cliente:
        headers.append("Valor Cliente")
    if ctx.show_tarifa_tercero:
        headers.append("Valor Tercero")
    if ctx.show_cointra_financials:
        headers.extend(["Rentabilidad", "Ganancia Cointra"])
    headers.append("Observaciones")

    col_idx = _xl_write_headers(ws, start_row + 1, headers, ctx)
    row_num = start_row + 2
    total_cliente = 0.0
    total_tercero = 0.0
    sorted_items = sorted(
        source_items,
        key=lambda current: (
            str(current.placa or "").strip().upper(),
            current.fecha_servicio,
            current.id,
        ),
    )

    if not sorted_items:
        ws.cell(row=row_num, column=1, value="(sin servicios)")
        ws.cell(row=row_num, column=1).font = Font(italic=True, color="6B7280")
        return row_num + 1, col_idx, total_cliente, total_tercero

    if write_plate_totals:
        grouped_items: dict[str, list[ConciliacionItem]] = {}
        for item in sorted_items:
            placa = str(item.placa or "").strip().upper() or "SIN_PLACA"
            grouped_items.setdefault(placa, []).append(item)

        ordered_placas = sorted(grouped_items.keys())
        for placa in ordered_placas:
            placa_total_cliente = 0.0
            placa_total_tercero = 0.0
            for item in grouped_items[placa]:
                tarifa_cliente = _as_float(item.tarifa_cliente)
                tarifa_tercero = _as_float(item.tarifa_tercero)
                total_cliente += tarifa_cliente
                total_tercero += tarifa_tercero
                placa_total_cliente += tarifa_cliente
                placa_total_tercero += tarifa_tercero
                values = {
                    "Placa": placa,
                    "Tipo Vehiculo": ctx.tipo_vehiculo_by_placa.get(placa, ""),
                    "Fecha": str(item.fecha_servicio or ""),
                    "Titulo Servicio": item.viaje.titulo if item.viaje else "",
                    "Tipo Servicio": (item.servicio_nombre or "").strip()
                    or (item.servicio_codigo or "").strip()
                    or str(getattr(item.tipo, "value", item.tipo)),
                    "Observaciones": item.descripcion or "",
                }
                for header, column in col_idx.items():
                    if header in values:
                        ws.cell(row=row_num, column=column, value=values[header])
                _xl_write_financials(ws, row_num, col_idx, tarifa_cliente, tarifa_tercero, ctx)
                _xl_style_row(ws, row_num, len(headers), ctx)
                row_num += 1

            ws.cell(row=row_num, column=col_idx["Tipo Vehiculo"], value=f"{totals_label_prefix} {placa}")
            _xl_write_financials(ws, row_num, col_idx, placa_total_cliente, placa_total_tercero, ctx)
            _xl_style_row(ws, row_num, len(headers), ctx, fill=ctx.section_fill, font=ctx.section_font)
            row_num += 2

        return row_num, col_idx, total_cliente, total_tercero

    for item in sorted_items:
        placa = str(item.placa or "").strip().upper() or "SIN_PLACA"
        tarifa_cliente = _as_float(item.tarifa_cliente)
        tarifa_tercero = _as_float(item.tarifa_tercero)
        total_cliente += tarifa_cliente
        total_tercero += tarifa_tercero
        values = {
            "Placa": placa,
            "Tipo Vehiculo": ctx.tipo_vehiculo_by_placa.get(placa, ""),
            "Fecha": str(item.fecha_servicio or ""),
            "Titulo Servicio": item.viaje.titulo if item.viaje else "",
            "Tipo Servicio": (item.servicio_nombre or "").strip()
            or (item.servicio_codigo or "").strip()
            or str(getattr(item.tipo, "value", item.tipo)),
            "Observaciones": item.descripcion or "",
        }
        for header, column in col_idx.items():
            if header in values:
                ws.cell(row=row_num, column=column, value=values[header])
        _xl_write_financials(ws, row_num, col_idx, tarifa_cliente, tarifa_tercero, ctx)
        _xl_style_row(ws, row_num, len(headers), ctx)
        row_num += 1

    return row_num, col_idx, total_cliente, total_tercero
