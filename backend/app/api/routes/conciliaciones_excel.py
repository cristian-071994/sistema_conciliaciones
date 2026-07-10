from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from app.models.conciliacion import Conciliacion
from app.models.conciliacion_item import ConciliacionItem
from app.models.enums import ItemTipo, UserRole
from app.models.viaje import Viaje
from .conciliaciones_helpers import (
    _as_float,
    _display_estado,
    _extract_liquidacion_metadata,
    _item_servicio_codigo,
    _split_total_evenly,
)
from .conciliaciones_excel_renders import (
    _ExcelContext,
    _xl_write_report_header,
    _xl_write_headers,
    _xl_style_row,
    _xl_write_financials,
    _xl_accumulate_totals,
    _xl_manifest_context,
    _xl_item_service_code_upper,
    _xl_item_service_label,
    _xl_write_adicionales_unified_section,
    _xl_write_transport_section,
    _xl_write_additional_services_section,
)


def _build_conciliacion_excel(
    conc: Conciliacion,
    items: list[ConciliacionItem],
    user_role: UserRole,
    tipo_vehiculo_by_placa: dict[str, str],
    avansat_prefetched: dict[str, dict] | None = None,
) -> bytes:
    liquidacion_items: list[tuple[ConciliacionItem, dict]] = []
    non_liquidacion_items: list[ConciliacionItem] = []
    for item in items:
        liq_meta = _extract_liquidacion_metadata(item)
        if liq_meta:
            liquidacion_items.append((item, liq_meta))
        else:
            non_liquidacion_items.append(item)

    has_liquidacion = bool(liquidacion_items)

    liquidacion_placas = {
        str(item.placa or "").strip().upper()
        for item, _ in liquidacion_items
        if str(item.placa or "").strip()
    }

    quincena_items: list[ConciliacionItem] = []
    additional_items: list[ConciliacionItem] = []
    for item in non_liquidacion_items:
        placa = str(item.placa or "").strip().upper()
        service_code = _xl_item_service_code_upper(item)
        if service_code == "DISPONIBILIDAD":
            continue
        if has_liquidacion:
            if item.tipo == ItemTipo.VIAJE and service_code == "VIAJE" and placa in liquidacion_placas:
                quincena_items.append(item)
            else:
                additional_items.append(item)
        else:
            if service_code == "VIAJE":
                quincena_items.append(item)
            else:
                additional_items.append(item)

    wb = Workbook()
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    ws_quincena = wb.create_sheet("Quincena")
    ws_adicionales = wb.create_sheet("Adicionales")

    header_fill = PatternFill(fill_type="solid", fgColor="E5E7EB")
    header_font = Font(bold=True, color="1F2937")
    section_fill = PatternFill(fill_type="solid", fgColor="FDE68A")
    section_font = Font(bold=True, color="111827")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    cop_format = '"$" #,##0'
    pct_format = '#,##0.##" %"'

    show_tarifa_tercero = user_role != UserRole.CLIENTE
    show_tarifa_cliente = user_role != UserRole.TERCERO
    show_cointra_financials = user_role == UserRole.COINTRA

    estado_display = _display_estado(conc)
    estado_label = estado_display.replace("_", " ")
    estado_styles = {
        "BORRADOR": ("ECFDF5", "065F46"),
        "EN_REVISION": ("FFFBEB", "92400E"),
        "APROBADA": ("F0FDFA", "115E59"),
        "ENVIADA_A_FACTURAR": ("F0F9FF", "075985"),
    }
    estado_fill_color, estado_font_color = estado_styles.get(estado_display, ("F7FEE7", "3F6212"))

    ctx = _ExcelContext(
        conc=conc,
        avansat_prefetched=avansat_prefetched,
        show_tarifa_tercero=show_tarifa_tercero,
        show_tarifa_cliente=show_tarifa_cliente,
        show_cointra_financials=show_cointra_financials,
        border=border,
        cop_format=cop_format,
        pct_format=pct_format,
        header_fill=header_fill,
        header_font=header_font,
        section_fill=section_fill,
        section_font=section_font,
        center=center,
        tipo_vehiculo_by_placa=tipo_vehiculo_by_placa,
        estado_label=estado_label,
        estado_fill_color=estado_fill_color,
        estado_font_color=estado_font_color,
    )

    liquidacion_items_sorted = sorted(
        liquidacion_items,
        key=lambda pair: (
            str(pair[0].placa or "").strip().upper(),
            pair[0].fecha_servicio,
            pair[0].id,
        ),
    )
    liquidacion_totals_by_placa = _xl_accumulate_totals([item for item, _ in liquidacion_items_sorted])

    if has_liquidacion:
        current_row = _xl_write_report_header(ws_resumen, "Resumen", ctx)
        ws_resumen.cell(row=current_row, column=1, value="LIQUIDACION CONTRATO FIJO")
        ws_resumen.cell(row=current_row, column=1).font = Font(bold=True, size=11, color="111827")
        top_headers = ["Placa", "Tipo Vehiculo"]
        if show_tarifa_cliente:
            top_headers.append("Valor Cliente")
        if show_tarifa_tercero:
            top_headers.append("Valor Tercero")
        if show_cointra_financials:
            top_headers.extend(["Rentabilidad", "Ganancia Cointra"])
        top_col_idx = _xl_write_headers(ws_resumen, current_row + 1, top_headers, ctx)
        current_row += 2

        top_total_cliente = 0.0
        top_total_tercero = 0.0
        for item, _ in liquidacion_items_sorted:
            placa = str(item.placa or "").strip().upper()
            tarifa_cliente = _as_float(item.tarifa_cliente)
            tarifa_tercero = _as_float(item.tarifa_tercero)
            top_total_cliente += tarifa_cliente
            top_total_tercero += tarifa_tercero
            ws_resumen.cell(row=current_row, column=top_col_idx["Placa"], value=placa)
            ws_resumen.cell(
                row=current_row,
                column=top_col_idx["Tipo Vehiculo"],
                value=tipo_vehiculo_by_placa.get(placa, ""),
            )
            _xl_write_financials(ws_resumen, current_row, top_col_idx, tarifa_cliente, tarifa_tercero, ctx)
            _xl_style_row(ws_resumen, current_row, len(top_headers), ctx)
            current_row += 1

        ws_resumen.cell(row=current_row, column=2, value="TOTAL CONTRATO FIJO")
        _xl_write_financials(ws_resumen, current_row, top_col_idx, top_total_cliente, top_total_tercero, ctx)
        _xl_style_row(ws_resumen, current_row, len(top_headers), ctx, fill=section_fill, font=section_font)
        current_row += 2

        summary_row, summary_bottom_idx, summary_total_cliente, summary_total_tercero = _xl_write_additional_services_section(
            ws_resumen,
            current_row,
            "ADICIONALES",
            additional_items,
            ctx,
        )
        if additional_items:
            ws_resumen.cell(row=summary_row, column=2, value="TOTAL ADICIONALES")
            _xl_write_financials(ws_resumen, summary_row, summary_bottom_idx, summary_total_cliente, summary_total_tercero, ctx)
            _xl_style_row(
                ws_resumen,
                summary_row,
                len(summary_bottom_idx),
                ctx,
                fill=section_fill,
                font=section_font,
            )

        current_row = _xl_write_report_header(ws_quincena, "Quincena", ctx)
        current_row, quincena_col_idx, _, _, quincena_totals_by_placa = _xl_write_transport_section(
            ws_quincena,
            current_row,
            "VIAJES",
            quincena_items,
            ctx,
        )
        current_row += 2

        ws_quincena.cell(row=current_row, column=1, value="DISPONIBILIDAD")
        ws_quincena.cell(row=current_row, column=1).font = Font(bold=True, size=11, color="111827")
        disponibilidad_headers = ["Placa", "Tipo Vehiculo"]
        if show_tarifa_cliente:
            disponibilidad_headers.append("Valor Cliente")
        if show_tarifa_tercero:
            disponibilidad_headers.append("Valor Tercero")
        if show_cointra_financials:
            disponibilidad_headers.append("Ganancia Cointra")
        disponibilidad_col_idx = _xl_write_headers(ws_quincena, current_row + 1, disponibilidad_headers, ctx)
        current_row += 2

        for placa in sorted(liquidacion_totals_by_placa.keys()):
            liq_totals = liquidacion_totals_by_placa.get(placa, {"cliente": 0.0, "tercero": 0.0})
            viajes_totals = quincena_totals_by_placa.get(placa, {"cliente": 0.0, "tercero": 0.0})
            disponibilidad_cliente = liq_totals["cliente"] - viajes_totals["cliente"]
            disponibilidad_tercero = liq_totals["tercero"] - viajes_totals["tercero"]
            ws_quincena.cell(row=current_row, column=disponibilidad_col_idx["Placa"], value=placa)
            ws_quincena.cell(
                row=current_row,
                column=disponibilidad_col_idx["Tipo Vehiculo"],
                value=tipo_vehiculo_by_placa.get(placa, ""),
            )
            if "Valor Cliente" in disponibilidad_col_idx:
                c = ws_quincena.cell(
                    row=current_row,
                    column=disponibilidad_col_idx["Valor Cliente"],
                    value=disponibilidad_cliente,
                )
                c.number_format = cop_format
            if "Valor Tercero" in disponibilidad_col_idx:
                c = ws_quincena.cell(
                    row=current_row,
                    column=disponibilidad_col_idx["Valor Tercero"],
                    value=disponibilidad_tercero,
                )
                c.number_format = cop_format
            if "Ganancia Cointra" in disponibilidad_col_idx:
                c = ws_quincena.cell(
                    row=current_row,
                    column=disponibilidad_col_idx["Ganancia Cointra"],
                    value=disponibilidad_cliente - disponibilidad_tercero,
                )
                c.number_format = cop_format
            _xl_style_row(ws_quincena, current_row, len(disponibilidad_headers), ctx)
            current_row += 1

        current_row += 2
        ws_quincena.cell(row=current_row, column=1, value="TOTALES POR VEHICULO")
        ws_quincena.cell(row=current_row, column=1).font = Font(bold=True, size=11, color="111827")
        resumen_quincena_headers = ["Placa", "Concepto"]
        if show_tarifa_cliente:
            resumen_quincena_headers.append("Valor Cliente")
        if show_tarifa_tercero:
            resumen_quincena_headers.append("Valor Tercero")
        if show_cointra_financials:
            resumen_quincena_headers.append("Ganancia Cointra")
        resumen_quincena_col_idx = _xl_write_headers(ws_quincena, current_row + 1, resumen_quincena_headers, ctx)
        current_row += 2

        for placa in sorted(liquidacion_totals_by_placa.keys()):
            liq_totals = liquidacion_totals_by_placa.get(placa, {"cliente": 0.0, "tercero": 0.0})
            viajes_totals = quincena_totals_by_placa.get(placa, {"cliente": 0.0, "tercero": 0.0})
            disponibilidad_cliente = liq_totals["cliente"] - viajes_totals["cliente"]
            disponibilidad_tercero = liq_totals["tercero"] - viajes_totals["tercero"]
            rows = [
                ("TOTAL VIAJES", viajes_totals["cliente"], viajes_totals["tercero"]),
                ("TOTAL DISPONIBILIDAD", disponibilidad_cliente, disponibilidad_tercero),
                ("TOTAL LIQUIDADO", viajes_totals["cliente"] + disponibilidad_cliente, viajes_totals["tercero"] + disponibilidad_tercero),
            ]
            for label, total_cliente, total_tercero in rows:
                ws_quincena.cell(row=current_row, column=resumen_quincena_col_idx["Placa"], value=placa)
                ws_quincena.cell(row=current_row, column=resumen_quincena_col_idx["Concepto"], value=label)
                if "Valor Cliente" in resumen_quincena_col_idx:
                    c = ws_quincena.cell(
                        row=current_row,
                        column=resumen_quincena_col_idx["Valor Cliente"],
                        value=total_cliente,
                    )
                    c.number_format = cop_format
                if "Valor Tercero" in resumen_quincena_col_idx:
                    c = ws_quincena.cell(
                        row=current_row,
                        column=resumen_quincena_col_idx["Valor Tercero"],
                        value=total_tercero,
                    )
                    c.number_format = cop_format
                if "Ganancia Cointra" in resumen_quincena_col_idx:
                    c = ws_quincena.cell(
                        row=current_row,
                        column=resumen_quincena_col_idx["Ganancia Cointra"],
                        value=total_cliente - total_tercero,
                    )
                    c.number_format = cop_format
                _xl_style_row(ws_quincena, current_row, len(resumen_quincena_headers), ctx, font=section_font)
                current_row += 1
            current_row += 1

        current_row += 1
        ws_quincena.cell(row=current_row, column=1, value="CONSOLIDADO FACTURA")
        ws_quincena.cell(row=current_row, column=1).font = Font(bold=True, size=11, color="111827")
        consolidated_headers = ["Concepto"]
        if show_tarifa_cliente:
            consolidated_headers.append("Valor Cliente")
        if show_tarifa_tercero:
            consolidated_headers.append("Valor Tercero")
        if show_cointra_financials:
            consolidated_headers.append("Ganancia Cointra")
        consolidated_col_idx = _xl_write_headers(ws_quincena, current_row + 1, consolidated_headers, ctx)
        current_row += 2

        total_viajes_cliente = sum(values.get("cliente", 0.0) for values in quincena_totals_by_placa.values())
        total_viajes_tercero = sum(values.get("tercero", 0.0) for values in quincena_totals_by_placa.values())
        total_disponibilidad_cliente = 0.0
        total_disponibilidad_tercero = 0.0
        for placa in liquidacion_totals_by_placa.keys():
            liq_totals = liquidacion_totals_by_placa.get(placa, {"cliente": 0.0, "tercero": 0.0})
            viajes_totals = quincena_totals_by_placa.get(placa, {"cliente": 0.0, "tercero": 0.0})
            total_disponibilidad_cliente += liq_totals["cliente"] - viajes_totals["cliente"]
            total_disponibilidad_tercero += liq_totals["tercero"] - viajes_totals["tercero"]

        consolidated_rows = [
            ("FACTURA VIAJES", total_viajes_cliente, total_viajes_tercero),
            ("FACTURA DISPONIBILIDAD", total_disponibilidad_cliente, total_disponibilidad_tercero),
        ]
        for label, total_cliente, total_tercero in consolidated_rows:
            ws_quincena.cell(row=current_row, column=consolidated_col_idx["Concepto"], value=label)
            if "Valor Cliente" in consolidated_col_idx:
                c = ws_quincena.cell(
                    row=current_row,
                    column=consolidated_col_idx["Valor Cliente"],
                    value=total_cliente,
                )
                c.number_format = cop_format
            if "Valor Tercero" in consolidated_col_idx:
                c = ws_quincena.cell(
                    row=current_row,
                    column=consolidated_col_idx["Valor Tercero"],
                    value=total_tercero,
                )
                c.number_format = cop_format
            if "Ganancia Cointra" in consolidated_col_idx:
                c = ws_quincena.cell(
                    row=current_row,
                    column=consolidated_col_idx["Ganancia Cointra"],
                    value=total_cliente - total_tercero,
                )
                c.number_format = cop_format
            _xl_style_row(ws_quincena, current_row, len(consolidated_headers), ctx, fill=section_fill, font=section_font)
            current_row += 1

        current_row += 1
        factura_total_cliente = total_viajes_cliente + total_disponibilidad_cliente
        factura_total_tercero = total_viajes_tercero + total_disponibilidad_tercero
        ws_quincena.cell(row=current_row, column=consolidated_col_idx["Concepto"], value="TOTAL FACTURA")
        if "Valor Cliente" in consolidated_col_idx:
            c = ws_quincena.cell(
                row=current_row,
                column=consolidated_col_idx["Valor Cliente"],
                value=factura_total_cliente,
            )
            c.number_format = cop_format
        if "Valor Tercero" in consolidated_col_idx:
            c = ws_quincena.cell(
                row=current_row,
                column=consolidated_col_idx["Valor Tercero"],
                value=factura_total_tercero,
            )
            c.number_format = cop_format
        if "Ganancia Cointra" in consolidated_col_idx:
            c = ws_quincena.cell(
                row=current_row,
                column=consolidated_col_idx["Ganancia Cointra"],
                value=factura_total_cliente - factura_total_tercero,
            )
            c.number_format = cop_format
        _xl_style_row(ws_quincena, current_row, len(consolidated_headers), ctx, fill=section_fill, font=section_font)
    else:
        current_row = _xl_write_report_header(ws_resumen, "Resumen", ctx)
        summary_row, summary_viajes_idx, summary_viajes_cliente, summary_viajes_tercero = _xl_write_additional_services_section(
            ws_resumen,
            current_row,
            "VIAJES",
            quincena_items,
            ctx,
            write_plate_totals=True,
            totals_label_prefix="TOTAL VEHICULO",
        )
        ws_resumen.cell(row=summary_row, column=2, value="TOTAL VIAJES")
        _xl_write_financials(ws_resumen, summary_row, summary_viajes_idx, summary_viajes_cliente, summary_viajes_tercero, ctx)
        _xl_style_row(
            ws_resumen,
            summary_row,
            len(summary_viajes_idx),
            ctx,
            fill=section_fill,
            font=section_font,
        )

        if additional_items:
            current_row = summary_row + 2
            summary_row, summary_bottom_idx, summary_total_cliente, summary_total_tercero = _xl_write_additional_services_section(
                ws_resumen,
                current_row,
                "ADICIONALES",
                additional_items,
                ctx,
                write_plate_totals=True,
                totals_label_prefix="TOTAL VEHICULO",
            )
            ws_resumen.cell(row=summary_row, column=2, value="TOTAL ADICIONALES")
            _xl_write_financials(ws_resumen, summary_row, summary_bottom_idx, summary_total_cliente, summary_total_tercero, ctx)
            _xl_style_row(
                ws_resumen,
                summary_row,
                len(summary_bottom_idx),
                ctx,
                fill=section_fill,
                font=section_font,
            )

        current_row = _xl_write_report_header(ws_quincena, "Quincena", ctx)
        current_row, quincena_col_idx, total_viajes_cliente, total_viajes_tercero, _ = _xl_write_transport_section(
            ws_quincena,
            current_row,
            "VIAJES",
            quincena_items,
            ctx,
        )

        current_row += 2
        consolidated_headers = ["Concepto"]
        if show_tarifa_cliente:
            consolidated_headers.append("Valor Cliente")
        if show_tarifa_tercero:
            consolidated_headers.append("Valor Tercero")
        if show_cointra_financials:
            consolidated_headers.append("Ganancia Cointra")
        consolidated_col_idx = _xl_write_headers(ws_quincena, current_row, consolidated_headers, ctx)
        current_row += 1
        ws_quincena.cell(row=current_row, column=consolidated_col_idx["Concepto"], value="TOTAL VIAJES")
        if "Valor Cliente" in consolidated_col_idx:
            c = ws_quincena.cell(row=current_row, column=consolidated_col_idx["Valor Cliente"], value=total_viajes_cliente)
            c.number_format = cop_format
        if "Valor Tercero" in consolidated_col_idx:
            c = ws_quincena.cell(row=current_row, column=consolidated_col_idx["Valor Tercero"], value=total_viajes_tercero)
            c.number_format = cop_format
        if "Ganancia Cointra" in consolidated_col_idx:
            c = ws_quincena.cell(row=current_row, column=consolidated_col_idx["Ganancia Cointra"], value=total_viajes_cliente - total_viajes_tercero)
            c.number_format = cop_format
        _xl_style_row(ws_quincena, current_row, len(consolidated_headers), ctx, fill=section_fill, font=section_font)

    if additional_items:
        current_row = _xl_write_report_header(ws_adicionales, "Adicionales", ctx)
        current_row, _, additional_trip_total_cliente, additional_trip_total_tercero, additional_services_total_cliente, additional_services_total_tercero = _xl_write_adicionales_unified_section(
            ws_adicionales,
            current_row,
            "ADICIONALES",
            additional_items,
            ctx,
        )
        current_row += 2

        totals_headers = ["Concepto"]
        if show_tarifa_cliente:
            totals_headers.append("Valor Cliente")
        if show_tarifa_tercero:
            totals_headers.append("Valor Tercero")
        if show_cointra_financials:
            totals_headers.append("Ganancia Cointra")
        totals_col_idx = _xl_write_headers(ws_adicionales, current_row, totals_headers, ctx)
        current_row += 1
        totals_rows = [
            ("FACTURA VIAJES", additional_trip_total_cliente, additional_trip_total_tercero),
            ("FACTURA SERVICIOS", additional_services_total_cliente, additional_services_total_tercero),
        ]
        for label, total_cliente, total_tercero in totals_rows:
            ws_adicionales.cell(row=current_row, column=totals_col_idx["Concepto"], value=label)
            if "Valor Cliente" in totals_col_idx:
                c = ws_adicionales.cell(row=current_row, column=totals_col_idx["Valor Cliente"], value=total_cliente)
                c.number_format = cop_format
            if "Valor Tercero" in totals_col_idx:
                c = ws_adicionales.cell(row=current_row, column=totals_col_idx["Valor Tercero"], value=total_tercero)
                c.number_format = cop_format
            if "Ganancia Cointra" in totals_col_idx:
                c = ws_adicionales.cell(row=current_row, column=totals_col_idx["Ganancia Cointra"], value=total_cliente - total_tercero)
                c.number_format = cop_format
            _xl_style_row(ws_adicionales, current_row, len(totals_headers), ctx, fill=section_fill, font=section_font)
            current_row += 1

        current_row += 1
        total_adic_cliente = additional_trip_total_cliente + additional_services_total_cliente
        total_adic_tercero = additional_trip_total_tercero + additional_services_total_tercero
        ws_adicionales.cell(row=current_row, column=totals_col_idx["Concepto"], value="TOTAL ADICIONALES")
        if "Valor Cliente" in totals_col_idx:
            c = ws_adicionales.cell(row=current_row, column=totals_col_idx["Valor Cliente"], value=total_adic_cliente)
            c.number_format = cop_format
        if "Valor Tercero" in totals_col_idx:
            c = ws_adicionales.cell(row=current_row, column=totals_col_idx["Valor Tercero"], value=total_adic_tercero)
            c.number_format = cop_format
        if "Ganancia Cointra" in totals_col_idx:
            c = ws_adicionales.cell(row=current_row, column=totals_col_idx["Ganancia Cointra"], value=total_adic_cliente - total_adic_tercero)
            c.number_format = cop_format
        _xl_style_row(ws_adicionales, current_row, len(totals_headers), ctx, fill=section_fill, font=section_font)
    else:
        wb.remove(ws_adicionales)

    resumen_widths = {
        "Placa": 16,
        "Tipo Vehiculo": 22,
        "Fecha": 16,
        "Titulo Servicio": 28,
        "Tipo Servicio": 24,
        "Valor Cliente": 18,
        "Valor Tercero": 18,
        "Rentabilidad": 16,
        "Ganancia Cointra": 20,
        "Observaciones": 36,
    }
    quincena_widths = {
        "Manifiesto": 18,
        "Fecha Emision": 18,
        "Placa Vehiculo": 16,
        "Trayler": 16,
        "Remesa": 18,
        "Producto": 28,
        "Ciudad Origen": 22,
        "Ciudad Destino": 22,
        "Valor Cliente": 18,
        "Valor Tercero": 18,
        "Rentabilidad": 16,
        "Ganancia Cointra": 20,
        "Placa": 16,
        "Tipo Vehiculo": 22,
        "Concepto": 24,
    }
    adicionales_widths = {**quincena_widths, **resumen_widths}

    sheets_with_widths = [
        (ws_resumen, resumen_widths),
        (ws_quincena, quincena_widths),
    ]
    if "Adicionales" in wb.sheetnames:
        sheets_with_widths.append((wb["Adicionales"], adicionales_widths))

    for sheet, widths in sheets_with_widths:
        min_column_width = 14
        max_column = sheet.max_column
        for col_idx in range(1, max_column + 1):
            header_value = str(sheet.cell(row=1, column=col_idx).value or "")
            best_width = widths.get(header_value, 18)
            for row_idx in range(1, min(sheet.max_row, 12) + 1):
                header_candidate = str(sheet.cell(row=row_idx, column=col_idx).value or "")
                if header_candidate in widths:
                    best_width = widths[header_candidate]
                    break
            sheet.column_dimensions[get_column_letter(col_idx)].width = max(best_width, min_column_width)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
